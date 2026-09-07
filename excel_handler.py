"""
Excel Handler — Upload parsing and formatted report generation.
"""

import io
import logging
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from platform_config import get_column_mapping, PLATFORM_CONFIG

logger = logging.getLogger("osa_qc.excel")


# ---------------------------------------------------------------------------
# Upload parsing
# ---------------------------------------------------------------------------

def parse_upload(
    uploaded_file,
    platform: str,
    sheet_name: str | int = 0,
) -> pd.DataFrame:
    """
    Parse an uploaded Excel file and return a DataFrame.

    The file is read as-is. The column mapping from platform_config is
    attached as metadata so the engine knows which column holds the URL,
    SKU ID, etc.

    Args:
        uploaded_file: Streamlit UploadedFile or file-like object
        platform: Platform key to apply column mapping
        sheet_name: Sheet name or index to read (default: first sheet)

    Returns:
        pd.DataFrame with column_mapping and platform in attrs
    """
    try:
        df = pd.read_excel(uploaded_file, sheet_name=sheet_name, engine="openpyxl")
        logger.info(f"Parsed upload: {len(df)} rows, columns: {list(df.columns)}")

        col_map = get_column_mapping(platform)
        df.attrs["column_mapping"] = col_map
        df.attrs["platform"] = platform

        # Validate that required columns exist
        missing = []
        for logical_name, actual_col in col_map.items():
            if actual_col not in df.columns:
                missing.append(f"{logical_name} (expected column '{actual_col}')")

        if missing:
            logger.warning(f"Missing columns in upload: {missing}")

        return df

    except Exception as e:
        logger.error(f"Failed to parse uploaded file: {e}")
        raise ValueError(f"Could not read the uploaded file: {e}")


def get_sheet_names(uploaded_file) -> list[str]:
    """Get sheet names from an uploaded Excel file."""
    try:
        xl = pd.ExcelFile(uploaded_file, engine="openpyxl")
        return xl.sheet_names
    except Exception:
        return []


# ---------------------------------------------------------------------------
# Report generation
# ---------------------------------------------------------------------------

# Styling constants
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
MISMATCH_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
ERROR_FILL = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
IN_STOCK_FONT = Font(name="Calibri", color="38761D", bold=True)
OOS_FONT = Font(name="Calibri", color="CC0000", bold=True)
CELL_FONT = Font(name="Calibri", size=10)
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)


def generate_report(
    results: dict[str, pd.DataFrame],
    include_summary: bool = True,
) -> io.BytesIO:
    """
    Generate a formatted Excel report from QC mismatch results.

    Args:
        results: Dict mapping platform names to DataFrames of mismatched rows.
                 Each DataFrame should have 'Actual_Status' and 'Remark' columns.
        include_summary: If True and multiple platforms, add a combined summary sheet.

    Returns:
        BytesIO buffer containing the .xlsx file, ready for download.
    """
    output = io.BytesIO()
    wb = Workbook()

    # Remove default sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    # If multiple platforms and summary requested, create combined sheet first
    if include_summary and len(results) > 1:
        all_mismatches = []
        for platform, df in results.items():
            if not df.empty:
                df_copy = df.copy()
                df_copy.insert(0, "Platform", PLATFORM_CONFIG.get(platform, {}).get(
                    "display_name", platform.title()
                ))
                all_mismatches.append(df_copy)

        if all_mismatches:
            combined = pd.concat(all_mismatches, ignore_index=True)
            _write_sheet(wb, "All Mismatches", combined)

    # One sheet per platform
    for platform, df in results.items():
        display_name = PLATFORM_CONFIG.get(platform, {}).get(
            "display_name", platform.title()
        )
        # Sheet names max 31 chars
        sheet_name = display_name[:31]
        if df.empty:
            ws = wb.create_sheet(title=sheet_name)
            ws["A1"] = "No mismatches found"
            ws["A1"].font = Font(name="Calibri", size=12, bold=True, color="38761D")
        else:
            _write_sheet(wb, sheet_name, df)

    # If no sheets were created (empty results), add a placeholder
    if not wb.sheetnames:
        ws = wb.create_sheet(title="Results")
        ws["A1"] = "No mismatches found across any platform"
        ws["A1"].font = Font(name="Calibri", size=12, bold=True, color="38761D")

    wb.save(output)
    output.seek(0)
    logger.info(f"Generated report with {len(results)} platform sheet(s)")
    return output


def _write_sheet(wb: Workbook, sheet_name: str, df: pd.DataFrame) -> None:
    """Write a DataFrame to a formatted worksheet."""
    ws = wb.create_sheet(title=sheet_name)

    # Write headers
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=str(col_name))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    # Write data rows
    for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
        for col_idx, (col_name, value) in enumerate(row.items(), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

            # Color-code the Remark column
            if col_name == "Remark":
                if "could not verify" in str(value).lower() or "error" in str(value).lower():
                    cell.fill = ERROR_FILL
                else:
                    cell.fill = MISMATCH_FILL

            # Color-code status columns
            if col_name in ("Actual_Status", "Remark"):
                val_lower = str(value).lower()
                if "in stock" in val_lower and "out" not in val_lower:
                    cell.font = IN_STOCK_FONT
                elif "out of stock" in val_lower or "oos" in val_lower:
                    cell.font = OOS_FONT

    # Auto-fit column widths (approximate)
    for col_idx in range(1, len(df.columns) + 1):
        col_letter = get_column_letter(col_idx)
        max_width = len(str(df.columns[col_idx - 1])) + 2

        for row_idx in range(2, min(len(df) + 2, 102)):  # Sample first 100 rows
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value:
                max_width = max(max_width, min(len(str(cell_value)), 50))

        ws.column_dimensions[col_letter].width = max_width + 2

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = ws.dimensions


def get_report_filename(platforms: list[str]) -> str:
    """Generate a descriptive filename for the report."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if len(platforms) == 1:
        platform_str = platforms[0]
    elif len(platforms) <= 3:
        platform_str = "_".join(platforms)
    else:
        platform_str = "multi_platform"
    return f"OSA_QC_Mismatches_{platform_str}_{timestamp}.xlsx"
